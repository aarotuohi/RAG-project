"""
Excel cost table builder — produces a laskentapohja-formatted .xlsx file
from generated section 2 data and project metadata.

Column layout (matching laskentapohja.xlsx):
  A — step index (narrow)
  B — name / label (wide)
  C — hourly rate  "88.0€/h"
  D — hours        "10 h"
  E — price        (number)
"""
from __future__ import annotations
from pathlib import Path
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from backend.config import OUTPUTS_DIR
from backend.chains.extraction_chain import ProjectData
from backend.ingestion.excel_parser import CostStepGroup


# ── Style helpers ─────────────────────────────────────────────────────────────

_BOLD = Font(bold=True)
_NORMAL = Font(bold=False)
_PHASE_FILL = PatternFill("solid", fgColor="D9E1F2")   # light blue — phase header
_TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")   # light grey — subtotal rows
_SUMMARY_FILL = PatternFill("solid", fgColor="BDD7EE")  # medium blue — summary
_INTERNAL_FILL = PatternFill("solid", fgColor="EFEFEF")  # very light grey — internal calc (fixed-price)


def _thin_border(bottom_only: bool = False) -> Border:
    thin = Side(style="thin")
    if bottom_only:
        return Border(bottom=thin)
    return Border(top=thin, bottom=thin, left=thin, right=thin)


def _set_cell(ws, row: int, col: int, value, bold: bool = False,
              fill: PatternFill | None = None, number_format: str | None = None,
              wrap: bool = False, align: str = "left"):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=bold)
    cell.alignment = Alignment(horizontal=align, vertical="top", wrap_text=wrap)
    if fill:
        cell.fill = fill
    if number_format:
        cell.number_format = number_format
    return cell


def _parse_project_datetime(value: str) -> datetime:
    if not value:
        return datetime.today()
    try:
        return datetime.fromisoformat(value)
    except ValueError:
        pass
    try:
        day, month, year = (int(part) for part in value.split("."))
        return datetime(year, month, day)
    except Exception:
        return datetime.today()


# ── Main builder ─────────────────────────────────────────────────────────────

# ── Translations ─────────────────────────────────────────────────────────────

_STRINGS: dict[str, dict[str, str]] = {
    "fi": {
        "sheet_title":      "Laskenta",
        "offer_no":         "Tarjous nro",
        "name":             "Nimi",
        "customer":         "Asiakas",
        "sales_resp":       "Myyntivastuu",
        "date":             "Päiväys",
        "description":      "Kuvaus",
        "personnel":        "Henkilöt",
        "phase":            "Vaihe",
        "col_rate":         "Tuntihinta [€/h]",
        "col_hours":        "Tuntiarvio [h]",
        "col_price":        "Hinta-arvio [€]",
        "persons_note":     "hlö",
        "phase_subtotal":   "Arvioidut työkustannukset yhteensä",
        "fixed_price_label": "Kiinteä hinta",
        "summary_header":   "Yhteensä",
        "summary_col_h":    "Tuntiarvio [h]",
        "summary_col_e":    "Hinta-arvio [€]",
        "summary_col_fp":   "Kiinteä hinta [€]",
        "tax_select_label": "Sisällytä ALV 25,5%",
        "tax_no":           "Ei",
        "tax_yes":          "Kyllä",
        "tax_amount_label": "ALV (25,5%)",
        "tax_total_label":  "Hinta yhteensä (sis. ALV)",
    },
    "en": {
        "sheet_title":      "Calculation",
        "offer_no":         "Offer no",
        "name":             "Name",
        "customer":         "Customer",
        "sales_resp":       "Sales responsible",
        "date":             "Date",
        "description":      "Description",
        "personnel":        "Personnel",
        "phase":            "Phase",
        "col_rate":         "Hourly rate [€/h]",
        "col_hours":        "Hours estimate [h]",
        "col_price":        "Price estimate [€]",
        "persons_note":     "pers",
        "phase_subtotal":   "Estimated labour costs total",
        "fixed_price_label": "Fixed price",
        "summary_header":   "Total",
        "summary_col_h":    "Hours estimate [h]",
        "summary_col_e":    "Price estimate [€]",
        "summary_col_fp":   "Fixed price [€]",
        "tax_select_label": "Include VAT 25.5%",
        "tax_no":           "No",
        "tax_yes":          "Yes",
        "tax_amount_label": "VAT (25.5%)",
        "tax_total_label":  "Total price (incl. VAT)",
    },
}


def build_cost_excel(project: ProjectData, section2: dict,
                     language: str = "fi") -> Path:
    """
    Build a laskentapohja-style Excel workbook from project info and section2
    cost data.  Saves to OUTPUTS_DIR and returns the file path.

    For fixed-price projects the sub-step rate/hours rows are rendered as
    internal (grey, italic) calculation rows and each phase shows a single
    fixed-price row instead of the hourly subtotal.
    """
    t = _STRINGS.get(language, _STRINGS["fi"])
    is_fixed = (section2.get("payment_type") or "").lower() == "fixed"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = t["sheet_title"]

    # Column widths (matching the template)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 85
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 18

    steps: list[CostStepGroup] = section2.get("steps", [])
    grand_total: float = section2.get("grand_total", 0.0)

    _d = _parse_project_datetime(project.document_date)
    _date_fi = f"{_d.day}.{_d.month}.{_d.year}"

    # ── Metadata header (rows 1-7) ────────────────────────────────────────────
    meta_rows = [
        (t["offer_no"],    project.project_number or ""),
        (t["name"],        project.project_name or ""),
        (t["customer"],    _customer_label(project)),
        (t["sales_resp"],  project.salesperson_name or ""),
        (t["date"],        _date_fi),
        (t["description"], section2.get("short_description") or project.goals or ""),
        (t["personnel"],   project.required_expertise or ""),
    ]
    for i, (label, value) in enumerate(meta_rows, start=1):
        _set_cell(ws, i, 2, label, bold=True, align="right")
        _set_cell(ws, i, 3, value, wrap=True)

    current_row = len(meta_rows) + 2  # blank separator row

    # ── Phase blocks ──────────────────────────────────────────────────────────
    phase_subtotal_rows: list[int] = []   # E-col subtotal row per phase (for summary)

    for phase_num, sg in enumerate(steps, start=1):
        phase_label = f"{t['phase']} {phase_num}. {sg.name}"

        # Phase header row
        _set_cell(ws, current_row, 2, phase_label, bold=True, fill=_PHASE_FILL)
        if is_fixed:
            # Fixed-price: only show the fixed-price column header
            _set_cell(ws, current_row, 5, t["fixed_price_label"], bold=True,
                      fill=_PHASE_FILL, align="right")
        else:
            _set_cell(ws, current_row, 3, t["col_rate"],  bold=True,
                      fill=_PHASE_FILL, align="right")
            _set_cell(ws, current_row, 4, t["col_hours"], bold=True,
                      fill=_PHASE_FILL, align="right")
            _set_cell(ws, current_row, 5, t["col_price"], bold=True,
                      fill=_PHASE_FILL, align="right")
        current_row += 1

        # Sub-step rows — numeric C & D so formulas work when user edits them
        sub_step_start = current_row
        for idx, ss in enumerate(sg.sub_steps, start=1):
            effective_hours = round(ss.hours * ss.persons, 1)
            name = ss.name + (f"  (×{ss.persons} {t['persons_note']})" if ss.persons > 1 else "")
            if is_fixed:
                # Render as internal/grey calculation rows (rate × hours still
                # calculated so the phase fixed price can be a rounded version)
                _set_cell(ws, current_row, 1, idx, align="center", fill=_INTERNAL_FILL)
                cell_name = ws.cell(row=current_row, column=2, value=name)
                cell_name.font = Font(italic=True, color="808080")
                cell_name.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell_name.fill = _INTERNAL_FILL
                _set_cell(ws, current_row, 3, ss.hourly_rate,
                          number_format='#,##0 \\\u20ac', align="right", fill=_INTERNAL_FILL)
                _set_cell(ws, current_row, 4, effective_hours,
                          number_format='#,##0.0 \\h', align="right", fill=_INTERNAL_FILL)
                _set_cell(ws, current_row, 5, f"=C{current_row}*D{current_row}",
                          number_format='#,##0.00', align="right", fill=_INTERNAL_FILL)
            else:
                _set_cell(ws, current_row, 1, idx, align="center")
                _set_cell(ws, current_row, 2, name, wrap=True)
                # C: numeric hourly rate — user can edit; drives column E
                _set_cell(ws, current_row, 3, ss.hourly_rate,
                          number_format='#,##0 \\\u20ac', align="right")
                # D: numeric hours — user can edit; drives column E and subtotal
                _set_cell(ws, current_row, 4, effective_hours,
                          number_format='#,##0.0 \\h', align="right")
                # E: live formula = rate × hours
                _set_cell(ws, current_row, 5, f"=C{current_row}*D{current_row}",
                          number_format='#,##0.00', align="right")
            current_row += 1
        sub_step_end = current_row - 1

        # Phase subtotal / fixed-price row
        if is_fixed:
            phase_cost = round(sum(ss.hourly_rate * ss.hours * ss.persons for ss in sg.sub_steps), 2)
            _set_cell(ws, current_row, 2, t["fixed_price_label"],
                      bold=True, fill=_TOTAL_FILL)
            # Fixed price = calculated cost (user can override the number directly)
            _set_cell(ws, current_row, 5, phase_cost,
                      bold=True, fill=_TOTAL_FILL,
                      number_format='#,##0.00', align="right")
        else:
            _set_cell(ws, current_row, 2, t["phase_subtotal"],
                      bold=True, fill=_TOTAL_FILL)
            _set_cell(ws, current_row, 4,
                      f"=SUM(D{sub_step_start}:D{sub_step_end})",
                      bold=True, fill=_TOTAL_FILL,
                      number_format='#,##0.0 \\h', align="right")
            _set_cell(ws, current_row, 5,
                      f"=SUM(E{sub_step_start}:E{sub_step_end})",
                      bold=True, fill=_TOTAL_FILL,
                      number_format='#,##0.00', align="right")
        phase_subtotal_rows.append(current_row)
        current_row += 1

        # Output row
        if sg.output:
            _set_cell(ws, current_row, 2, f"Output: {sg.output}", wrap=True)
            current_row += 1

        current_row += 1  # blank row between phases

    # ── Summary section ───────────────────────────────────────────────────────
    _set_cell(ws, current_row, 2, t["summary_header"],   bold=True, fill=_SUMMARY_FILL)
    if is_fixed:
        _set_cell(ws, current_row, 5, t["summary_col_fp"], bold=True, fill=_SUMMARY_FILL, align="right")
    else:
        _set_cell(ws, current_row, 4, t["summary_col_h"],    bold=True, fill=_SUMMARY_FILL, align="right")
        _set_cell(ws, current_row, 5, t["summary_col_e"],    bold=True, fill=_SUMMARY_FILL, align="right")
    current_row += 1

    summary_rows: list[int] = []
    for i, (sg, st_row) in enumerate(zip(steps, phase_subtotal_rows), start=1):
        _set_cell(ws, current_row, 1, i, align="center")
        _set_cell(ws, current_row, 2, f"{t['phase']} {i}: {sg.name}", wrap=True)
        if is_fixed:
            # Reference the fixed-price cell directly (no hours column in summary)
            _set_cell(ws, current_row, 5, f"=E{st_row}",
                      number_format='#,##0.00', align="right")
        else:
            _set_cell(ws, current_row, 4, f"=D{st_row}",
                      number_format='#,##0.0', align="right")
            _set_cell(ws, current_row, 5, f"=E{st_row}",
                      number_format='#,##0.00', align="right")
        summary_rows.append(current_row)
        current_row += 1

    # Grand total row — SUM over all summary rows
    grand_total_row = current_row
    if summary_rows:
        if not is_fixed:
            _set_cell(ws, current_row, 4,
                      f"=SUM(D{summary_rows[0]}:D{summary_rows[-1]})",
                      bold=True, fill=_SUMMARY_FILL,
                      number_format='#,##0.0', align="right")
        _set_cell(ws, current_row, 5,
                  f"=SUM(E{summary_rows[0]}:E{summary_rows[-1]})",
                  bold=True, fill=_SUMMARY_FILL,
                  number_format='#,##0.00', align="right")
    else:
        if not is_fixed:
            _set_cell(ws, current_row, 4, 0,
                      bold=True, fill=_SUMMARY_FILL,
                      number_format='#,##0.0', align="right")
        _set_cell(ws, current_row, 5, 0,
                  bold=True, fill=_SUMMARY_FILL,
                  number_format='#,##0.00', align="right")
    current_row += 1

    # ── VAT / Tax section ─────────────────────────────────────────────────────
    current_row += 1  # blank separator row

    # Tax-selection row: label in B, dropdown in C
    tax_select_row = current_row
    _set_cell(ws, current_row, 2, t["tax_select_label"], bold=True)
    tax_cell = ws.cell(row=current_row, column=3, value=t["tax_no"])
    tax_cell.font = Font(bold=True)
    tax_cell.alignment = Alignment(horizontal="center", vertical="top")
    # Dropdown: "Ei" / "Kyllä"  (or "No" / "Yes" in English)
    dv = DataValidation(
        type="list",
        formula1=f'"{t["tax_no"]},{t["tax_yes"]}"',
        allow_blank=False,
        showDropDown=False,  # False = show the arrow button in Excel
    )
    ws.add_data_validation(dv)
    dv.add(tax_cell)
    current_row += 1

    # VAT amount row
    vat_amount_row = current_row
    _set_cell(ws, current_row, 2, t["tax_amount_label"])
    _set_cell(ws, current_row, 5,
              f'=IF(C{tax_select_row}="{t["tax_yes"]}",E{grand_total_row}*0.255,0)',
              number_format='#,##0.00', align="right")
    current_row += 1

    # Final total row (net + VAT)
    _TAX_TOTAL_FILL = PatternFill("solid", fgColor="D6E4BC")   # light green — tax total
    _set_cell(ws, current_row, 2, t["tax_total_label"], bold=True, fill=_TAX_TOTAL_FILL)
    _set_cell(ws, current_row, 5,
              f"=E{grand_total_row}+E{vat_amount_row}",
              bold=True, fill=_TAX_TOTAL_FILL,
              number_format='#,##0.00', align="right")
    current_row += 1

    
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1   
    ws.page_setup.fitToHeight = 0 
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:E{current_row}"
    ws.print_title_rows = "1:7"   

    # ── Save ──────────────────────────────────────────────────────────────────
    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    stem = (project.project_name or "offer").replace(" ", "_")[:40]
    out_path = OUTPUTS_DIR / f"{stem}_cost_table.xlsx"

    # Avoid overwriting
    counter = 1
    while out_path.exists():
        out_path = OUTPUTS_DIR / f"{stem}_cost_table_{counter}.xlsx"
        counter += 1

    wb.save(str(out_path))
    return out_path


def _customer_label(project: ProjectData) -> str:
    """Build a display name for the customer field."""
    name = project.customer_name
    if project.company_name:
        return f"{project.company_name}{(' / ' + name) if name else ''}"
    return name
