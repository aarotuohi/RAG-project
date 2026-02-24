"""
Document loader — dispatches files to the correct parser and indexes them into ChromaDB.
Supported formats: .docx, .pptx, .xlsx, .xls, .txt, .md, .pdf
"""
from pathlib import Path
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document

from backend.config import CHUNK_SIZE, CHUNK_OVERLAP
from backend.vectorstore.chroma_client import get_collection


def load_file(file_path: Path) -> list[Document]:
    """Load a file and return a list of LangChain Document objects."""
    suffix = file_path.suffix.lower()

    if suffix in (".docx",):
        return _load_docx(file_path)
    elif suffix in (".pptx",):
        return _load_pptx(file_path)
    elif suffix in (".xlsx", ".xls"):
        return _load_excel_as_text(file_path)
    elif suffix in (".txt", ".md"):
        return _load_text(file_path)
    elif suffix == ".pdf":
        return _load_pdf(file_path)
    else:
        return []


def _load_docx(path: Path) -> list[Document]:
    from docx import Document as DocxDocument
    doc = DocxDocument(str(path))
    full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [Document(page_content=full_text, metadata={"source": str(path), "type": "docx"})]


def _load_pptx(path: Path) -> list[Document]:
    from pptx import Presentation
    prs = Presentation(str(path))
    docs = []
    for i, slide in enumerate(prs.slides):
        text = "\n".join(
            shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()
        )
        if text.strip():
            docs.append(Document(page_content=text, metadata={"source": str(path), "slide": i + 1, "type": "pptx"}))
    return docs


def _load_excel_as_text(path: Path) -> list[Document]:
    import pandas as pd
    docs = []
    xl = pd.ExcelFile(str(path))
    for sheet in xl.sheet_names:
        df = xl.parse(sheet).fillna("")
        text = f"Sheet: {sheet}\n" + df.to_string(index=False)
        docs.append(Document(page_content=text, metadata={"source": str(path), "sheet": sheet, "type": "xlsx"}))
    return docs


def _load_text(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [Document(page_content=text, metadata={"source": str(path), "type": "txt"})]


def _load_pdf(path: Path) -> list[Document]:
    import pdfplumber
    docs = []
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            if text.strip():
                docs.append(Document(page_content=text, metadata={"source": str(path), "page": i + 1, "type": "pdf"}))
    return docs


def index_file(file_path: Path, collection_name: str, extra_metadata: dict | None = None) -> int:
    """
    Load, chunk, and index a file into the specified ChromaDB collection.
    Returns the number of chunks added.
    """
    raw_docs = load_file(file_path)
    if not raw_docs:
        return 0

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
    )
    chunks = splitter.split_documents(raw_docs)

    if extra_metadata:
        for chunk in chunks:
            chunk.metadata.update(extra_metadata)

    collection = get_collection(collection_name)
    collection.add_documents(chunks)
    return len(chunks)


def index_directory(directory: Path, collection_name: str, extra_metadata: dict | None = None) -> dict:
    """Index all supported files in a directory. Returns {filename: chunk_count}."""
    results = {}
    supported = {".docx", ".pptx", ".xlsx", ".xls", ".txt", ".md", ".pdf"}
    for f in directory.iterdir():
        if f.is_file() and f.suffix.lower() in supported:
            count = index_file(f, collection_name, extra_metadata)
            results[f.name] = count
    return results
